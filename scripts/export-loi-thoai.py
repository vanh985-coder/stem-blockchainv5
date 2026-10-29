"""Bước 2 của xuất Excel lời thoại: đọc JSON do scripts/export-loi-thoai.test.mjs tạo, ghi docs/LOI-THOAI-NHAN-VAT.xlsx.

Dùng: python3 scripts/export-loi-thoai.py <json-vào> <xlsx-ra>   (thường gọi qua `node scripts/export-loi-thoai.mjs`)
Chỉ định dạng và trình bày; không tự thêm hay sửa lời thoại.
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

src, dst = sys.argv[1], sys.argv[2]
data = json.load(open(src, encoding='utf-8'))

HEADER_FILL = PatternFill('solid', fgColor='4A3B8C')
HEADER_FONT = Font(bold=True, color='FFFFFF', name='Calibri', size=11)
BODY_FONT = Font(name='Calibri', size=11)
THIN = Side(style='thin', color='D0CCE0')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(wrap_text=True, vertical='top')
CENTER_TOP = Alignment(wrap_text=True, vertical='top', horizontal='center')


def write_table(ws, headers, rows, widths, center_cols=()):
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.border = BORDER
            if cell.row == 1:
                cell.font = HEADER_FONT
                cell.fill = HEADER_FILL
                cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
            else:
                cell.font = BODY_FONT
                cell.alignment = CENTER_TOP if cell.column in center_cols else WRAP_TOP
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = 'C2' if len(headers) > 8 else 'B2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{ws.max_row}'


DIALOGUE_HEADERS = [
    'STT', 'Nhân vật', 'Làng/khu vực', 'Số màn', 'Tên màn', 'Cảnh/trạm/độ khó', 'Điều kiện xuất hiện',
    'Thứ tự trong đoạn hội thoại', 'Nội dung nguyên văn', 'Khóa nội dung', 'File nguồn và nơi sử dụng', 'Ghi chú',
]
DIALOGUE_WIDTHS = [6, 16, 24, 8, 26, 34, 52, 14, 72, 38, 62, 48]


def dialogue_row(r):
    nguon = r['src']
    if r.get('also'):
        nguon += '\n\nCâu này cũng xuất hiện ở:\n' + r['also']
    return [
        r['stt'], r['nv'], r['lang'], r['man'] if r['man'] else None, r['manTen'], r['canh'], r['dk'],
        r['thuTu'], r['text'], r['key'], nguon, r['note'],
    ]


wb = Workbook()
wb.properties.title = 'Lời thoại nhân vật - Giấc mơ Sổ Chung'
wb.properties.creator = 'scripts/export-loi-thoai'

# 1. Lời thoại (đang xuất hiện trong game)
ws = wb.active
ws.title = 'Lời thoại'
write_table(ws, DIALOGUE_HEADERS, [dialogue_row(r) for r in data['rows']], DIALOGUE_WIDTHS, center_cols=(1, 4, 8))

# 2. Nhân vật
ws = wb.create_sheet('Nhân vật')
rows = []
for c in data['characters']:
    if c['id'] == 'hocSinh':
        ghi = 'Người chơi. Không có tên cố định, không có lượt thoại riêng: lời thoại gọi học sinh bằng {ten} / {Ten}.'
    elif c['id'] == 'phanDien':
        ghi = 'MUỐN ĐỔI TÊN CHỈ SỬA hằng số phanDien ở content/characters.ts. Mọi lời thoại viết {phanDien}, không viết thẳng tên.'
    elif c['id'] == 'laiBuon':
        ghi = 'Có chân dung và model nhưng chưa có lời thoại nào trong game.'
    elif c['id'] == 'coChi':
        ghi = 'Chỉ nói 1 câu (Bài 2, trước Trạm Khó).'
    else:
        ghi = ''
    rows.append([c['id'], c['name'], c['label'], c['placeholder'], c['portrait'], c['count'], c['places'] or '(chưa có lời)', ghi])
rows.append(['(không phải nhân vật)', 'Lời dẫn truyện', 'Lời dẫn truyện', '', '', sum(1 for r in data['rows'] if r['nv'] == 'Lời dẫn truyện'), 'Câu chuyện mở đầu, kết thúc, mở đầu mỗi làng, cao trào', 'Văn bản kể chuyện hiện cùng ảnh truyện (story.ts). Có lúc chứa lời Bi, phản diện hay cô giáo trong ngoặc kép.'])
rows.append(['(không phải nhân vật)', 'Biển gỗ', 'Biển gỗ', '', '', 0, 'Xem sheet "Không rõ người nói"', 'Hộp thoại ghi tên người nói là "Biển gỗ" khi chưa đủ thẻ "Em có biết?".'])
write_table(
    ws,
    ['Mã nhân vật (CharacterId)', 'Tên hiện tại (content/characters.ts)', 'Tên ở hộp thoại', 'Chỗ giữ tên trong lời thoại', 'File chân dung (ui/portraits/)', 'Số lượt thoại (sheet Lời thoại)', 'Xuất hiện ở', 'Ghi chú'],
    rows,
    [24, 30, 22, 26, 26, 16, 60, 70],
    center_cols=(6,),
)

# 3. Biến
ws = wb.create_sheet('Biến')
rows = []
for v in data['variables']:
    used = ', '.join(str(x) for x in v['usedIn'])
    name = data['names'].get(v['key'])
    rows.append(['{' + v['key'] + '}', name or '', v['meaning'], used or '(không dùng trong sheet Lời thoại)'])
write_table(ws, ['Biến', 'Giá trị hiện tại (nếu là tên nhân vật)', 'Ý nghĩa', 'STT các dòng dùng biến trong sheet Lời thoại'], rows, [18, 28, 80, 60])

# 4. Chưa dùng
ws = wb.create_sheet('Chưa dùng')
write_table(ws, DIALOGUE_HEADERS, [dialogue_row(r) for r in data['unused']], DIALOGUE_WIDTHS, center_cols=(1, 4, 8))

# 5. Không rõ người nói
ws = wb.create_sheet('Không rõ người nói')
write_table(ws, DIALOGUE_HEADERS, [dialogue_row(r) for r in data['unclear']], DIALOGUE_WIDTHS, center_cols=(1, 4, 8))

# 6. Phạm vi
ws = wb.create_sheet('Phạm vi')
c = data['counts']
chars_with_lines = [x for x in data['characters'] if x['count'] > 0]
scope = [
    ['Tổng số lượt thoại (sheet Lời thoại)', f"{c['rows']} dòng, mỗi dòng một lượt nói; trong đó {c['narration']} dòng là Lời dẫn truyện (không phải nhân vật)."],
    ['Số nhân vật có lời', f"{len(chars_with_lines)} nhân vật có chân dung: " + ', '.join(f"{x['label']} ({x['count']})" for x in chars_with_lines) + '. Học sinh (người chơi) và lái buôn không có lời. Ngoài ra có 2 "người nói" không phải nhân vật: Lời dẫn truyện, Biển gỗ.'],
    ['Nội dung chưa dùng', f"{c['unused']} dòng ở sheet Chưa dùng (cổng đóng, hội làng đã ra mắt, hiệu ứng mở cổng, biển sắp ra mắt)."],
    ['Không rõ người nói', f"{c['unclear']} nhóm ở sheet Không rõ người nói."],
    ['Cách đếm', 'Mỗi dòng sheet Lời thoại là một lượt nói trong một tình huống. Cùng một câu dùng ở nhiều nơi (ví dụ lời trao Trang Sổ Vàng ở trang bài học và ở cảnh làng; ảnh truyện 6–11 ở /truyen và ở cảnh làng/cao trào) được đếm thành nhiều dòng, mỗi dòng ghi "Câu này cũng xuất hiện ở" trong cột File nguồn.'],
    ['Được tính là lời thoại', 'Chữ do một nhân vật nói trong hộp thoại có tên/chân dung (VnDialog, DialogueBox), bong bóng lời Bi ở màn 3, hộp Bi ở màn 12, "Mẹo từ Bi" ở thẻ giới thiệu trạm, bảng Hết tim có chân dung Bi, bong bóng "nghĩ thầm" của phản diện, câu phản diện nói ở Bài 4, và các lượt Lời dẫn truyện của ảnh truyện (đánh dấu rõ).'],
    ['Không tính', 'Tên nút, nhãn HUD, câu hỏi và đáp án, lời giải thích đáp án (FeedbackSheet), câu khen sau đáp án đúng, thẻ "Em có biết?" đọc trong màn (không có người nói), Hướng dẫn chơi, lời nhắc điều khiển, nhãn đọc màn hình, mô tả mục tiêu trạm, banner kết quả.'],
    ['Thẻ "Em có biết?" của người dẫn', 'Chỉ tính khi người dẫn đọc ở lượt giới thiệu làng trên /ban-do (VillageIntro). Thẻ nhặt trong làng và thẻ trên trang bài học không có người nói nên không đưa vào.'],
    ['Phạm vi chưa xác định 1', 'Câu kết thúc: ảnh truyện 12–14 chỉ chiếu trong trang /truyen; chưa có cảnh kết thúc/hội làng riêng trong game (hội làng chưa ra mắt, FESTIVAL_LAUNCHED = false).'],
    ['Phạm vi chưa xác định 2', 'Làng Bạc không có cảnh trao Trang Sổ Vàng ở làng và bài học 10 không có award: Trang Sổ Vàng thứ tư chỉ trao trong cảnh cao trào sau màn 12.'],
    ['Phạm vi chưa xác định 3', 'Cột "Điều kiện xuất hiện" tóm tắt theo mã hiện tại; các cờ "lần đầu" (đã xem cảnh, đã xem mẹo) lưu theo tài khoản hoặc cookie chơi thử nên một số lời chỉ hiện một lần.'],
    ['Phạm vi chưa xác định 4', 'Hai dòng có người nói nhập nhằng đã xếp ở sheet Lời thoại kèm ghi chú: lời {phanDien} hiện trong bong bóng chân dung Bi (màn 3), và chú thích "thầy Linh trao Trang Sổ Vàng thứ tư" (có thể là lời dẫn).'],
    ['Phạm vi chưa xác định 5', 'Các màn game 2, 5, 6, 8, 9, 11 không có lời nhân vật trong mã hiện tại (đã rà games.ts và các trang game): không có dòng nào ở sheet Lời thoại.'],
    ['Cách tạo lại file', 'Chạy `node scripts/export-loi-thoai.mjs` (cần vitest và Python 3 có openpyxl). Script chỉ đọc content/ và mã gọi; không sửa lời thoại hay logic game. Cần cập nhật bảng siêu dữ liệu (điều kiện, màn, thứ tự) trong scripts/export-loi-thoai.test.mjs khi thêm lời mới.'],
]
write_table(ws, ['Mục', 'Nội dung'], scope, [34, 150])
ws.freeze_panes = 'A2'

wb.save(dst)
print(f"Đã ghi {dst}: {c['rows']} lượt thoại")
